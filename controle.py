#!/usr/bin/env python3
"""Atualiza políticas (.docx) a partir de uma planilha de controle.

Para cada arquivo listado na planilha:
  1. Rodapé            -> data de revisão e versão
  2. Primeira página   -> data (Mês/Ano) da caixa de texto
  3. Final do documento -> nova linha no quadro "Registro de alterações"
                           (preenche a primeira linha vazia do quadro; se não
                           houver, acrescenta uma nova ao final)

Colunas obrigatórias: ARQUIVO, NOVA_VERSAO, DATA_REVISAO
Colunas opcionais:    ITEM_MODIFICADO, MODIFICACAO, MOTIVO, MES_ANO
                      (MES_ANO vazio/ausente = derivado de DATA_REVISAO)

Se o mesmo ARQUIVO aparecer em várias linhas, todas viram linhas no quadro
e o rodapé usa a de maior versão/data.
"""
from __future__ import annotations

import argparse
import copy
import logging
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

log = logging.getLogger("politicas")

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
         "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

COLUNAS_OBRIGATORIAS = ["ARQUIVO", "NOVA_VERSAO", "DATA_REVISAO"]
COLUNAS_QUADRO = ["VERSAO", "ITEM MODIFICADO", "MODIFICACAO", "MOTIVO", "DATA"]

RE_DATA = re.compile(r"\d{2}/\d{2}/\d{4}")
RE_VERSAO = re.compile(r"(\d+)\s*([ºª°]?)")
RE_MES_ANO = re.compile(rf"({'|'.join(MESES)})/\d{{4}}", re.IGNORECASE)

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


@dataclass(frozen=True)
class Linha:
    versao: int
    data: date
    mes_ano: str
    item: str
    modificacao: str
    motivo: str


# --------------------------------------------------------------------------
# Conversão dos valores da planilha
# --------------------------------------------------------------------------
def normalizar_versao(valor) -> int:
    if pd.isna(valor):
        raise ValueError("NOVA_VERSAO vazia")
    if isinstance(valor, str):
        valor = re.sub(r"[ºª°\s]", "", valor)
    try:
        numero = float(valor)
    except ValueError:
        raise ValueError(f"NOVA_VERSAO inválida: {valor!r}") from None
    if not numero.is_integer():
        raise ValueError(f"NOVA_VERSAO deve ser inteira: {valor!r}")
    return int(numero)


def normalizar_data(valor) -> date:
    if pd.isna(valor):
        raise ValueError("DATA_REVISAO vazia")
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ValueError(f"DATA_REVISAO inválida: {valor!r}")


def mes_ano_de(valor, data: date) -> str:
    padrao = f"{MESES[data.month - 1]}/{data.year}"
    if valor is None or pd.isna(valor):
        return padrao
    if isinstance(valor, date):  # datetime e Timestamp herdam de date
        return f"{MESES[valor.month - 1]}/{valor.year}"
    texto = str(valor).strip().capitalize() or padrao
    if not RE_MES_ANO.fullmatch(texto):
        raise ValueError(f"MES_ANO deve ser no formato Mês/AAAA: {valor!r}")
    return texto


def texto_ou_traco(valor) -> str:
    """Células em branco viram '-', como no padrão do quadro."""
    if valor is None or pd.isna(valor):
        return "-"
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip() or "-"


def converter_linha(idx: int, row: pd.Series) -> Linha:
    try:
        data = normalizar_data(row["DATA_REVISAO"])
        return Linha(
            versao=normalizar_versao(row["NOVA_VERSAO"]),
            data=data,
            mes_ano=mes_ano_de(row.get("MES_ANO"), data),
            item=texto_ou_traco(row.get("ITEM_MODIFICADO")),
            modificacao=texto_ou_traco(row.get("MODIFICACAO")),
            motivo=texto_ou_traco(row.get("MOTIVO")),
        )
    except ValueError as exc:
        raise ValueError(f"linha {idx + 2} da planilha: {exc}") from exc


# --------------------------------------------------------------------------
# Helpers de XML
# --------------------------------------------------------------------------
def texto_proprio(p) -> str:
    """Texto dos runs que são filhos diretos do parágrafo.

    Ignora parágrafos aninhados (caixas de texto), evitando contar duas vezes
    o conteúdo de mc:AlternateContent (Choice + Fallback).
    """
    return "".join(
        t.text or "" for r in p.findall(qn("w:r")) for t in r.findall(qn("w:t"))
    )


def tem_campo(p) -> bool:
    """True se o parágrafo contém campo do Word (PAGE, NUMPAGES, ...)."""
    return next(
        p.iter(qn("w:fldChar"), qn("w:instrText"), qn("w:fldSimple")), None
    ) is not None


def definir_texto(p, texto: str) -> None:
    """Troca o texto do parágrafo preservando a formatação do primeiro run.

    O Word fragmenta o texto em vários runs; aqui o texto novo vai inteiro
    no primeiro e os demais são esvaziados. Se o parágrafo não tem nenhum run
    com texto (célula vazia), cria um com a formatação do próprio parágrafo
    (pPr/rPr), que é onde o Word guarda a fonte de células vazias.
    """
    runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
    if not runs:
        run = OxmlElement("w:r")
        rpr_paragrafo = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
        if rpr_paragrafo is not None:
            run.append(copy.deepcopy(rpr_paragrafo))
        run.append(OxmlElement("w:t"))
        p.append(run)
        runs = [run]

    primeiro, *demais = runs
    t, *t_extras = primeiro.findall(qn("w:t"))
    t.text = texto
    t.set(XML_SPACE, "preserve")
    for extra in t_extras:
        primeiro.remove(extra)
    for run in demais:
        for t_antigo in run.findall(qn("w:t")):
            run.remove(t_antigo)
        if all(filho.tag == qn("w:rPr") for filho in run):
            p.remove(run)


def normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return " ".join(sem_acento.upper().split())


def texto_celula(tc) -> str:
    # split() também trata o espaço não separável (\xa0) que o Word costuma deixar
    return " ".join("".join(t.text or "" for t in tc.iter(qn("w:t"))).split())


# --------------------------------------------------------------------------
# 1) Rodapé: data de revisão e versão
# --------------------------------------------------------------------------
def rodapes_unicos(doc) -> list:
    """Elementos XML dos rodapés realmente definidos (sem repetir os 'vinculados')."""
    rodapes = []
    for secao in doc.sections:
        for rodape in (secao.footer, secao.first_page_footer, secao.even_page_footer):
            if rodape.is_linked_to_previous:
                continue
            elemento = rodape._element
            if not any(elemento is e for e in rodapes):
                rodapes.append(elemento)
    return rodapes


def campos_do_rodape(rodape):
    """Devolve (parágrafo da data de revisão, parágrafo da versão).

    O rodapé tem duas datas (criação e revisão, nessa ordem) e uma versão
    ("6º"). Parágrafos com campo (número da página) são ignorados.
    """
    candidatos = [p for p in rodape.iter(qn("w:p")) if not tem_campo(p)]
    datas = [p for p in candidatos if RE_DATA.fullmatch(texto_proprio(p).strip())]
    versoes = [p for p in candidatos if RE_VERSAO.fullmatch(texto_proprio(p).strip())]
    if len(datas) != 2 or len(versoes) != 1:
        raise ValueError(
            "layout do rodapé inesperado: esperava 2 datas (criação e revisão) "
            f"e 1 versão, encontrei {len(datas)} data(s) e {len(versoes)} versão(ões)"
        )
    return datas[1], versoes[0]


def atualizar_rodape(doc, revisao: str, versao: int) -> None:
    rodapes = rodapes_unicos(doc)
    if not rodapes:
        raise ValueError("documento sem rodapé")
    for rodape in rodapes:
        p_revisao, p_versao = campos_do_rodape(rodape)
        atual = RE_VERSAO.fullmatch(texto_proprio(p_versao).strip())
        if versao < int(atual.group(1)):
            log.warning("versão nova (%s) menor que a atual (%s)", versao, atual.group(1))
        definir_texto(p_revisao, revisao)
        definir_texto(p_versao, f"{versao}{atual.group(2)}")  # mantém o "º"


# --------------------------------------------------------------------------
# 2) Primeira página: Mês/Ano
# --------------------------------------------------------------------------
def paragrafos_mes_ano(doc):
    raizes = [doc.element.body]
    for secao in doc.sections:
        for cabecalho in (secao.header, secao.first_page_header, secao.even_page_header):
            if not cabecalho.is_linked_to_previous:
                raizes.append(cabecalho._element)
    for raiz in raizes:
        for p in raiz.iter(qn("w:p")):
            if RE_MES_ANO.fullmatch(texto_proprio(p).strip()):
                yield p


def atualizar_mes_ano(doc, mes_ano: str) -> None:
    paragrafos = list(paragrafos_mes_ano(doc))
    if not paragrafos:
        raise ValueError("data (Mês/Ano) da primeira página não encontrada")
    for p in paragrafos:
        definir_texto(p, mes_ano)


# --------------------------------------------------------------------------
# 3) Quadro "Registro de alterações"
# --------------------------------------------------------------------------
def localizar_quadro(doc):
    """Devolve (linhas <w:tr> do quadro, {coluna: índice}) pelo cabeçalho."""
    for tabela in doc.tables:
        linhas = tabela._tbl.findall(qn("w:tr"))
        if not linhas:
            continue
        cabecalho = [normalizar(texto_celula(tc)) for tc in linhas[0].findall(qn("w:tc"))]
        if all(coluna in cabecalho for coluna in COLUNAS_QUADRO):
            return linhas, {c: cabecalho.index(c) for c in COLUNAS_QUADRO}
    raise ValueError('quadro "Registro de alterações" não encontrado')


def linha_vazia(tr) -> bool:
    return all(texto_celula(tc) == "" for tc in tr.findall(qn("w:tc")))


def adicionar_linha_historico(doc, linha: Linha) -> bool:
    """Registra a alteração no quadro.

    O modelo do quadro traz linhas vazias já formatadas: usa a primeira delas.
    Só se não houver nenhuma, clona a última linha (mantém bordas e fonte).
    """
    linhas, idx = localizar_quadro(doc)
    valores = {
        "VERSAO": str(linha.versao),
        "ITEM MODIFICADO": linha.item,
        "MODIFICACAO": linha.modificacao,
        "MOTIVO": linha.motivo,
        "DATA": linha.data.strftime("%d/%m/%Y"),
    }

    for tr in linhas[1:]:  # idempotência: não duplica linha idêntica
        tcs = tr.findall(qn("w:tc"))
        if all(texto_celula(tcs[idx[c]]) == v for c, v in valores.items()):
            log.warning("linha já existe no quadro (versão %s), ignorada", linha.versao)
            return False

    destino = next((tr for tr in linhas[1:] if linha_vazia(tr)), None)
    if destino is None:
        modelo = linhas[-1]
        destino = copy.deepcopy(modelo)
        for elemento in destino.iter("*"):  # ids duplicados confundem o Word
            for atributo in ("paraId", "textId"):
                elemento.attrib.pop(f"{{{W14}}}{atributo}", None)
        modelo.addnext(destino)

    tcs = destino.findall(qn("w:tc"))
    for coluna, valor in valores.items():
        tc = tcs[idx[coluna]]
        paragrafos = tc.findall(qn("w:p"))
        for extra in paragrafos[1:]:
            tc.remove(extra)
        definir_texto(paragrafos[0], valor)
    return True


# --------------------------------------------------------------------------
# Orquestração
# --------------------------------------------------------------------------
def verificar_saida(caminho: Path, ref: Linha) -> None:
    """Reabre o arquivo gravado e confere cada alteração."""
    doc = Document(caminho)
    data = ref.data.strftime("%d/%m/%Y")

    for rodape in rodapes_unicos(doc):
        p_revisao, p_versao = campos_do_rodape(rodape)
        if texto_proprio(p_revisao).strip() != data:
            raise ValueError("verificação: data de revisão do rodapé não confere")
        if int(RE_VERSAO.fullmatch(texto_proprio(p_versao).strip()).group(1)) != ref.versao:
            raise ValueError("verificação: versão do rodapé não confere")

    capa = [texto_proprio(p).strip() for p in paragrafos_mes_ano(doc)]
    if not capa or any(t != ref.mes_ano for t in capa):
        raise ValueError("verificação: data da primeira página não confere")

    linhas, idx = localizar_quadro(doc)
    achou = any(
        texto_celula(tcs[idx["VERSAO"]]) == str(ref.versao)
        and texto_celula(tcs[idx["DATA"]]) == data
        for tcs in (tr.findall(qn("w:tc")) for tr in linhas[1:])
    )
    if not achou:
        raise ValueError("verificação: linha nova não encontrada no quadro")


def processar_arquivo(entrada: Path, saida: Path, linhas: list[Linha]) -> None:
    ref = max(linhas, key=lambda l: (l.versao, l.data))
    doc = Document(entrada)
    atualizar_rodape(doc, ref.data.strftime("%d/%m/%Y"), ref.versao)
    atualizar_mes_ano(doc, ref.mes_ano)
    novas = sum(adicionar_linha_historico(doc, linha) for linha in linhas)
    doc.save(saida)  # só grava se tudo acima funcionou
    verificar_saida(saida, ref)
    log.info("%s: versão %s, %s, %d linha(s) no quadro", saida.name, ref.versao, ref.mes_ano, novas)


def ler_planilha(caminho: Path) -> pd.DataFrame:
    df = pd.read_excel(caminho, engine="openpyxl")
    df.columns = [str(c).strip().upper() for c in df.columns]
    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise SystemExit(f"Colunas ausentes na planilha: {', '.join(faltando)}")
    return df.dropna(subset=["ARQUIVO"])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--planilha", type=Path, default=Path("Controle_Políticas.xlsx"))
    ap.add_argument("--entrada", type=Path, default=Path("Entrada"))
    ap.add_argument("--saida", type=Path, default=Path("Saida"))
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")
    args.saida.mkdir(parents=True, exist_ok=True)
    df = ler_planilha(args.planilha)

    ok = erros = 0
    for arquivo, grupo in df.groupby(df["ARQUIVO"].astype(str).str.strip(), sort=False):
        try:
            entrada = args.entrada / arquivo
            if not entrada.exists():
                raise FileNotFoundError("arquivo não encontrado na pasta de entrada")
            linhas = [converter_linha(i, r) for i, r in grupo.iterrows()]
            processar_arquivo(entrada, args.saida / arquivo, linhas)
            ok += 1
        except Exception as exc:  # um arquivo com erro não derruba o lote
            log.error("%s: %s", arquivo, exc)
            erros += 1

    log.info("Concluído: %d ok, %d com erro.", ok, erros)
    return 1 if erros else 0


if __name__ == "__main__":
    sys.exit(main())